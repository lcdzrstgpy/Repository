"""Excel 导入导出（契约第十四节）。

- 导出：直接返回 xlsx 文件流，**不走**统一 {code, msg, data} 包装（契约 14.1）
- 导入：multipart/form-data 上传 xlsx，逐行独立处理，返回统一响应结构（契约 14.2）

本模块的 router 不设 prefix：导出在 `/api/export/*`、导入在 `/api/import/*`，
两个前缀不同，故在各个路由装饰器里写全路径。
"""

import io
import logging
from datetime import datetime
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.response import BizException, ParamException, ok
from app.core.security import require_roles
from app.models.basic import Partner, Product, ProductSku, Warehouse
from app.models.inventory import Inventory
from app.models.order import SalesOrder, status_text
from app.models.user import SysUser
from app.schemas.serializers import fmt_dt

logger = logging.getLogger(__name__)

router = APIRouter(tags=["数据导入导出"])

# 导出文件 MIME（契约 14.1）
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# 单次导出行数上限：导出不分页，但要有上限保护，超出截断并记日志
MAX_EXPORT_ROWS = 10000
# 查询时多取一行，用于判断是否发生了截断
EXPORT_QUERY_LIMIT = MAX_EXPORT_ROWS + 1

# 通用状态中文名（契约 3.5）
GENERAL_STATUS_TEXT: dict[int, str] = {1: "启用", 0: "停用"}


# ---------------------------------------------------------------- 导出工具


def _status_label(value) -> str:
    """通用状态（1 启用 / 0 停用）转中文。"""
    if value is None:
        return ""
    return GENERAL_STATUS_TEXT.get(int(value), "")


def _new_workbook(title: str, headers: list[str], widths: list[int]):
    """新建工作簿：写表头（加粗）、设置列宽、冻结首行。"""
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(headers)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    return wb, ws


def _stream_xlsx(wb: Workbook, filename: str) -> StreamingResponse:
    """工作簿写入内存后以文件流返回。

    文件名是中文，按 RFC 5987 用 `filename*=UTF-8''<urlencoded>` 传递，
    否则浏览器下载时文件名会乱码（契约 14.1）。
    """
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type=XLSX_CONTENT_TYPE,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


def _cap(rows: list, label: str) -> list:
    """导出行数上限保护：超出 MAX_EXPORT_ROWS 时截断并记日志。"""
    if len(rows) > MAX_EXPORT_ROWS:
        logger.warning("%s导出结果超过 %d 行，已截断", label, MAX_EXPORT_ROWS)
        return rows[:MAX_EXPORT_ROWS]
    return rows


def _parse_date(raw: str | None, *, end_of_day: bool) -> datetime | None:
    """解析日期查询参数，支持 `YYYY-MM-DD` 与 `YYYY-MM-DD HH:MM:SS`。

    start_date 取当天 00:00:00，end_date 取当天 23:59:59（契约 14.1）。
    """
    text = (raw or "").strip()
    if not text:
        return None
    value: datetime | None = None
    date_only = False
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            value = datetime.strptime(text, fmt)
            date_only = fmt == "%Y-%m-%d"
            break
        except ValueError:
            continue
    if value is None:
        raise ParamException("日期格式不正确，应为 YYYY-MM-DD")
    if end_of_day and date_only:
        value = value.replace(hour=23, minute=59, second=59)
    return value


def _order_name_maps(db: Session, orders: list[SalesOrder]) -> tuple[dict, dict]:
    """批量取客户名与用户名，避免逐行查询。"""
    customer_ids = {o.customer_id for o in orders if o.customer_id}
    user_ids = {o.created_by for o in orders if o.created_by}
    user_ids |= {o.claimed_by for o in orders if o.claimed_by}

    customers = (
        {p.id: p.name for p in db.scalars(select(Partner).where(Partner.id.in_(customer_ids)))}
        if customer_ids
        else {}
    )
    users = (
        {
            u.id: (u.real_name or u.username)
            for u in db.scalars(select(SysUser).where(SysUser.id.in_(user_ids)))
        }
        if user_ids
        else {}
    )
    return customers, users


# ---------------------------------------------------------------- 导出接口


@router.get("/api/export/products", summary="导出商品 Excel")
def export_products(
    db: Session = Depends(get_db),
    _current_user: SysUser = Depends(require_roles("admin")),
):
    """商品：商品编码 / 商品名称 / 分类 / 单位 / 状态。仅 admin。"""
    rows = list(
        db.scalars(
            select(Product).order_by(Product.id.asc()).limit(EXPORT_QUERY_LIMIT)
        ).all()
    )
    rows = _cap(rows, "商品")

    wb, ws = _new_workbook(
        "商品", ["商品编码", "商品名称", "分类", "单位", "状态"], [18, 32, 16, 10, 10]
    )
    for product in rows:
        ws.append(
            [product.code, product.name, product.category, product.unit, _status_label(product.status)]
        )
    return _stream_xlsx(wb, "商品列表.xlsx")


@router.get("/api/export/skus", summary="导出 SKU Excel")
def export_skus(
    db: Session = Depends(get_db),
    _current_user: SysUser = Depends(require_roles("admin")),
):
    """SKU：SKU 编码 / 商品名称 / 规格 / 售价 / 状态。仅 admin。"""
    rows = list(
        db.execute(
            select(ProductSku, Product)
            .outerjoin(Product, Product.id == ProductSku.product_id)
            .order_by(ProductSku.id.asc())
            .limit(EXPORT_QUERY_LIMIT)
        ).all()
    )
    rows = _cap(rows, "SKU")

    wb, ws = _new_workbook(
        "SKU", ["SKU 编码", "商品名称", "规格", "售价", "状态"], [18, 32, 24, 12, 10]
    )
    for sku, product in rows:
        ws.append(
            [
                sku.sku_code,
                product.name if product else None,
                sku.spec,
                float(sku.price or 0),
                _status_label(sku.status),
            ]
        )
    return _stream_xlsx(wb, "SKU列表.xlsx")


@router.get("/api/export/inventory", summary="导出库存 Excel")
def export_inventory(
    warehouse_id: int | None = Query(None, description="按仓库筛选"),
    db: Session = Depends(get_db),
    _current_user: SysUser = Depends(require_roles("warehouse", "admin")),
):
    """库存：SKU 编码 / 商品名称 / 规格 / 仓库 / 库存量 / 预留量 / 可用量。

    仅 warehouse / admin；支持 `?warehouse_id=`（契约 14.1）。
    """
    stmt = (
        select(Inventory, ProductSku, Product, Warehouse)
        .outerjoin(ProductSku, ProductSku.id == Inventory.sku_id)
        .outerjoin(Product, Product.id == ProductSku.product_id)
        .outerjoin(Warehouse, Warehouse.id == Inventory.warehouse_id)
    )
    if warehouse_id is not None:
        stmt = stmt.where(Inventory.warehouse_id == warehouse_id)
    rows = list(
        db.execute(stmt.order_by(Inventory.id.asc()).limit(EXPORT_QUERY_LIMIT)).all()
    )
    rows = _cap(rows, "库存")

    wb, ws = _new_workbook(
        "库存",
        ["SKU 编码", "商品名称", "规格", "仓库", "库存量", "预留量", "可用量"],
        [18, 32, 24, 16, 12, 12, 12],
    )
    for inventory, sku, product, warehouse in rows:
        quantity = inventory.quantity or 0
        reserved = inventory.reserved_quantity or 0
        ws.append(
            [
                sku.sku_code if sku else None,
                product.name if product else None,
                sku.spec if sku else None,
                warehouse.name if warehouse else None,
                float(quantity),
                float(reserved),
                float(quantity - reserved),
            ]
        )
    return _stream_xlsx(wb, "库存列表.xlsx")


@router.get("/api/export/sales-orders", summary="导出销售订单 Excel")
def export_sales_orders(
    status: int | None = Query(None, description="订单状态"),
    start_date: str | None = Query(None, description="下单开始日期 YYYY-MM-DD"),
    end_date: str | None = Query(None, description="下单结束日期 YYYY-MM-DD"),
    db: Session = Depends(get_db),
    current_user: SysUser = Depends(require_roles("operator", "warehouse", "admin")),
):
    """订单：单号 / 客户 / 状态 / 总数量 / 总金额 / 物流单号 / 下单人 / 接单人 / 下单时间 / 发货时间。

    可见性规则与订单列表一致（契约 5.3 / 14.1）：operator 只导出自己创建的订单，
    warehouse / admin 导出全部。日期按 created_at 过滤，end_date 含当天 23:59:59。
    """
    start_at = _parse_date(start_date, end_of_day=False)
    end_at = _parse_date(end_date, end_of_day=True)

    conditions = []
    if current_user.role == "operator":
        conditions.append(SalesOrder.created_by == current_user.id)
    if status is not None:
        conditions.append(SalesOrder.status == status)
    if start_at is not None:
        conditions.append(SalesOrder.created_at >= start_at)
    if end_at is not None:
        conditions.append(SalesOrder.created_at <= end_at)

    orders = list(
        db.scalars(
            select(SalesOrder)
            .where(*conditions)
            .order_by(SalesOrder.id.asc())
            .limit(EXPORT_QUERY_LIMIT)
        ).all()
    )
    orders = _cap(orders, "销售订单")
    customers, users = _order_name_maps(db, orders)

    wb, ws = _new_workbook(
        "销售订单",
        [
            "单号",
            "客户",
            "状态",
            "总数量",
            "总金额",
            "物流单号",
            "下单人",
            "接单人",
            "下单时间",
            "发货时间",
        ],
        [22, 26, 12, 12, 14, 20, 12, 12, 22, 22],
    )
    for order in orders:
        ws.append(
            [
                order.no,
                customers.get(order.customer_id),
                status_text(order.status),
                float(order.total_count or 0),
                float(order.total_price or 0),
                order.express_no,
                users.get(order.created_by),
                users.get(order.claimed_by),
                fmt_dt(order.created_at),
                fmt_dt(order.shipped_at),
            ]
        )
    return _stream_xlsx(wb, "销售订单列表.xlsx")


# ---------------------------------------------------------------- 导入接口


def _cell_text(value) -> str:
    """单元格值转字符串：去空白；Excel 把数字编码读成 1.0 时还原成 1。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _cell_status(value) -> int | None:
    """解析状态列：接受 1 / 0（数字或字符串），空值默认 1；非法返回 None。"""
    if value is None:
        return 1
    text = str(value).strip()
    if not text:
        return 1
    if text in ("1", "1.0"):
        return 1
    if text in ("0", "0.0"):
        return 0
    return None


@router.post("/api/import/products", summary="导入商品 Excel")
def import_products(
    file: UploadFile = File(..., description="商品导入模板（.xlsx）"),
    db: Session = Depends(get_db),
    _current_user: SysUser = Depends(require_roles("admin")),
):
    """批量导入商品（契约 14.2），逐行独立处理，单行失败不影响其他行。仅 admin。

    模板列（第 1 行表头，第 2 行起数据）：
        商品编码 | 商品名称 | 分类 | 单位 | 状态
    错误信息里的 row 是 Excel 实际行号。
    """
    filename = (file.filename or "").strip()
    if not filename.lower().endswith(".xlsx"):
        raise BizException("仅支持 .xlsx 格式的文件")

    try:
        workbook = load_workbook(file.file, read_only=True, data_only=True)
    except Exception:
        raise BizException("文件无法解析，请确认上传的是有效的 .xlsx 文件")

    try:
        sheet = workbook.active
        # 已存在的商品编码一次性取出，避免逐行查询
        existing_codes = {code for code in db.scalars(select(Product.code)).all() if code}

        total = 0
        success = 0
        errors: list[dict] = []

        for row_no, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            # 整行空白（Excel 常见的尾部空行）直接跳过，不计入 total、不报错
            if all(_cell_text(cell) == "" for cell in row):
                continue

            total += 1
            code = _cell_text(row[0] if len(row) > 0 else None)
            name = _cell_text(row[1] if len(row) > 1 else None)
            category = _cell_text(row[2] if len(row) > 2 else None)
            unit = _cell_text(row[3] if len(row) > 3 else None)
            status = _cell_status(row[4] if len(row) > 4 else None)

            if not code:
                errors.append({"row": row_no, "message": "商品编码不能为空"})
                continue
            if not name:
                errors.append({"row": row_no, "message": "商品名称不能为空"})
                continue
            if status is None:
                errors.append(
                    {
                        "row": row_no,
                        "message": f"状态值不合法：{_cell_text(row[4])}，仅支持 1（启用）或 0（停用）",
                    }
                )
                continue
            if code in existing_codes:
                errors.append({"row": row_no, "message": f"商品编码已存在：{code}"})
                continue

            # 逐行提交：单行失败只回滚这一行，不影响已成功的行
            try:
                db.add(
                    Product(
                        code=code,
                        name=name,
                        category=category or None,
                        unit=unit or None,
                        status=status,
                    )
                )
                db.commit()
            except IntegrityError:
                db.rollback()
                errors.append({"row": row_no, "message": f"商品编码已存在：{code}"})
                continue
            except Exception as exc:  # noqa: BLE001 - 单行异常不阻断整批导入
                db.rollback()
                logger.warning("商品导入第 %d 行失败：%s", row_no, exc)
                errors.append({"row": row_no, "message": f"该行导入失败：{exc}"})
                continue

            existing_codes.add(code)
            success += 1
    finally:
        workbook.close()

    return ok(
        {
            "total": total,
            "success": success,
            "failed": len(errors),
            "errors": errors,
        },
        msg="导入完成",
    )
